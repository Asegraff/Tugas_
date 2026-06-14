"""Excel Workbook Management"""

import logging
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path
from config.settings import settings
from config.constants import EXCEL_SHEETS, RIGS, CATEGORIES

logger = logging.getLogger(__name__)

class ExcelWorkbook:
    """Manage Excel workbook operations"""
    
    def __init__(self, filepath: str = None):
        self.filepath = Path(filepath or settings.EXCEL_TEMPLATE_PATH)
        self.workbook = None
    
    def create_master_workbook(self) -> bool:
        """Create master workbook template"""
        try:
            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)
            
            # Create sheets
            self._create_dashboard_sheet()
            self._create_input_transaksi_sheet()
            self._create_master_rig_sheet()
            self._create_master_kategori_sheet()
            self._create_data_transaksi_sheet()
            self._create_invoice_generator_sheet()
            self._create_setting_sheet()
            self._create_log_sheet()
            
            self.workbook.save(self.filepath)
            logger.info(f"Master workbook created: {self.filepath}")
            return True
        except Exception as e:
            logger.error(f"Create workbook failed: {e}")
            return False
    
    def _create_dashboard_sheet(self):
        """Create Dashboard sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['DASHBOARD'], 0)
        self._style_title(ws, 'DASHBOARD')
        ws['A3'] = 'Ringkasan Sistem Invoice & Rekap ACU'
        ws.merge_cells('A3:D3')
    
    def _create_input_transaksi_sheet(self):
        """Create InputTransaksi sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['INPUT_TRANSAKSI'])
        headers = ['Tanggal', 'RIG', 'Kategori', 'Qty/Pax', 'Harga', 'Total']
        self._add_headers(ws, headers)
    
    def _create_master_rig_sheet(self):
        """Create MasterRig sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['MASTER_RIG'])
        headers = ['ID_Rig', 'NamaRig', 'Client']
        self._add_headers(ws, headers)
        
        # Add data
        row = 2
        for rig_id, rig_info in RIGS.items():
            ws[f'A{row}'] = rig_id
            ws[f'B{row}'] = rig_info['name']
            ws[f'C{row}'] = rig_info['client']
            row += 1
    
    def _create_master_kategori_sheet(self):
        """Create MasterKategori sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['MASTER_KATEGORI'])
        headers = ['ID_Kategori', 'NamaKategori']
        self._add_headers(ws, headers)
        
        # Add data
        row = 2
        for cat_id, cat_info in CATEGORIES.items():
            ws[f'A{row}'] = cat_id
            ws[f'B{row}'] = cat_info['name']
            row += 1
    
    def _create_data_transaksi_sheet(self):
        """Create DataTransaksi sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['DATA_TRANSAKSI'])
        headers = ['ID', 'Tanggal', 'Rig', 'Kategori', 'Qty_Pax', 'Harga', 'Total']
        self._add_headers(ws, headers)
    
    def _create_invoice_generator_sheet(self):
        """Create InvoiceGenerator sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['INVOICE_GENERATOR'])
        headers = ['ID_Invoice', 'NoInvoice', 'Rig', 'Periode', 'NilaiInvoice', 'Status', 'CreatedDate']
        self._add_headers(ws, headers)
    
    def _create_setting_sheet(self):
        """Create Setting sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['SETTING'])
        ws['A1'] = 'Parameter'
        ws['B1'] = 'Value'
        self._style_header(ws, ['A1', 'B1'])
        
        settings_data = [
            ('Database Path', settings.DB_PATH),
            ('PDF Output Path', settings.PDF_OUTPUT_PATH),
            ('Backup Path', settings.BACKUP_PATH),
            ('Backup Interval (days)', settings.BACKUP_INTERVAL),
        ]
        
        row = 2
        for param, value in settings_data:
            ws[f'A{row}'] = param
            ws[f'B{row}'] = str(value)
            row += 1
    
    def _create_log_sheet(self):
        """Create Log sheet"""
        ws = self.workbook.create_sheet(EXCEL_SHEETS['LOG'])
        headers = ['Timestamp', 'User', 'Action', 'Details']
        self._add_headers(ws, headers)
    
    def _add_headers(self, ws, headers):
        """Add header row with styling"""
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            self._style_header(ws, [cell.coordinate])
    
    def _style_header(self, ws, cells):
        """Apply header styling"""
        fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        font = Font(bold=True, color="FFFFFF")
        
        for cell_coord in cells:
            cell = ws[cell_coord]
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    def _style_title(self, ws, title):
        """Add title with styling"""
        ws['A1'] = title
        ws.merge_cells('A1:D1')
        cell = ws['A1']
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    def load_workbook(self) -> bool:
        """Load existing workbook"""
        try:
            if self.filepath.exists():
                self.workbook = load_workbook(self.filepath)
                logger.info(f"Workbook loaded: {self.filepath}")
                return True
            else:
                logger.warning(f"Workbook not found: {self.filepath}")
                return False
        except Exception as e:
            logger.error(f"Load workbook failed: {e}")
            return False
    
    def save(self) -> bool:
        """Save workbook"""
        try:
            if self.workbook:
                self.workbook.save(self.filepath)
                logger.info(f"Workbook saved: {self.filepath}")
                return True
            return False
        except Exception as e:
            logger.error(f"Save workbook failed: {e}")
            return False
